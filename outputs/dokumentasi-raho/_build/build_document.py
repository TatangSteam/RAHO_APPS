from pathlib import Path
import sys, math, textwrap, json
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[3]
OUT=Path(__file__).resolve().parents[1]
BUILD=Path(__file__).resolve().parent
SKILL=Path('C:/Users/jovan/.codex/plugins/cache/openai-primary-runtime/documents/26.601.10930/skills/documents')
sys.path.insert(0,str(SKILL/'scripts'))
from table_geometry import apply_table_geometry, audit_docx_tables
NAVY='#142B42'; BLUE='#2E74B5'; GOLD='#B58A32'; TEAL='#187D87'; LIGHT='#E8EEF5'; GRAY='#536475'
FONT='C:/Windows/Fonts/calibri.ttf'; BOLD='C:/Windows/Fonts/calibrib.ttf'

def font(n,bold=False): return ImageFont.truetype(BOLD if bold else FONT,n)
def canvas(title,subtitle='',height=640):
    im=Image.new('RGB',(1800,height),'white'); d=ImageDraw.Draw(im)
    d.text((35,22),title,font=font(43,True),fill=NAVY)
    if subtitle: d.text((35,79),subtitle,font=font(27),fill=GRAY)
    return im,d
def box(d,xy,title,detail='',fill=LIGHT):
    x,y,w,h=xy; d.rounded_rectangle((x,y,x+w,y+h),20,fill=fill,outline='#C5D2DD',width=2)
    lines=title.split('\n'); yy=y+22
    for line in lines:
        d.text((x+w/2,yy),line,font=font(32,True),fill=NAVY,anchor='mt'); yy+=39
    for line in detail.split('\n'):
        if line: d.text((x+w/2,yy+8),line,font=font(26),fill=GRAY,anchor='mt'); yy+=32
def arrow(d,a,b,color=TEAL):
    d.line((a,b),fill=color,width=5); ang=math.atan2(b[1]-a[1],b[0]-a[0]); l=18
    pts=[b,(b[0]-l*math.cos(ang-.5),b[1]-l*math.sin(ang-.5)),(b[0]-l*math.cos(ang+.5),b[1]-l*math.sin(ang+.5))]
    d.polygon(pts,fill=color)
def save(im,name): im.save(BUILD/(name+'.png')); return BUILD/(name+'.png')

im,d=canvas('Satu alur operasional, data saling terhubung','Peta proses utama RAHO ERP',580)
for x,t,s in [(35,'Member','Identitas & dokumen'),(480,'Paket & invoice','Harga, kuota, tagihan'),(925,'Pembayaran','Bukti & verifikasi'),(1370,'Sesi terapi','Plan & tindakan')]: box(d,(x,145,390,140),t,s)
for x in [425,870,1315]: arrow(d,(x,215),(x+55,215))
for x,t,s in [(1370,'Material & stok','Ledger / inventori tim'),(925,'Finance','HPP & jurnal'),(480,'Laporan','Operasional & insentif'),(35,'Integrasi','Zoho / WhatsApp')]: box(d,(x,365,390,140),t,s)
arrow(d,(1565,285),(1565,365))
for x in [1370,925,480]: arrow(d,(x,435),(x-55,435))
save(im,'flow_utama')

im,d=canvas('Arsitektur aplikasi','Komponen yang teridentifikasi pada repository',790)
box(d,(580,140,640,110),'Pengguna staff & member','Browser • akses sesuai role')
box(d,(580,315,640,120),'Web • Next.js 14 / React','TypeScript • halaman & komponen')
box(d,(580,500,640,115),'API • Express / Node.js','JWT • validasi • izin • service')
arrow(d,(900,250),(900,315));arrow(d,(900,435),(900,500))
box(d,(35,500,455,115),'PostgreSQL / Prisma','Data bisnis & transaksi');arrow(d,(580,557),(490,557))
box(d,(1310,500,455,115),'MinIO','Foto, dokumen & bukti');arrow(d,(1220,557),(1310,557))
box(d,(580,675,640,90),'Integrasi eksternal','Zoho Books • WhatsApp • RAIN')
arrow(d,(900,615),(900,675));save(im,'arsitektur')

im,d=canvas('Sesi terapi: pengisian sampai finalisasi','Pemilik proses adalah pembagian kerja; akses tetap mengikuti permission dan assignment',750)
items=[('Diagnosis','Dokter'),('Therapy plan','Dokter'),('Tanda vital sebelum','Nakes'),('Infus aktual','Nakes / Dokter'),('Material','Nakes'),('Foto / dokumentasi','Nakes'),('Tanda vital sesudah','Nakes'),('Evaluasi dokter','Dokter')]
for i,(t,s) in enumerate(items):
    row=i//4; col=i%4 if row==0 else 3-i%4; x=35+col*445;y=155+row*220
    box(d,(x,y,390,140),t,s)
    if i<3: arrow(d,(x+390,y+70),(x+445,y+70))
    if i==3: arrow(d,(x+195,y+140),(x+195,y+220))
    if i>3 and i<7: arrow(d,(x,y+70),(x-55,y+70))
d.rounded_rectangle((35,590,1760,705),18,fill=NAVY)
d.text((900,615),'Validasi kelengkapan → Selesaikan sesi → Stok, benefit & finance',font=font(32,True),fill='white',anchor='mt')
d.text((900,660),'Diagram merangkum komponen proses; bukan penomoran panel UI.',font=font(26),fill='#D6E3ED',anchor='mt');save(im,'flow_sesi')

im,d=canvas('Koreksi therapy plan yang sudah digunakan','Cabang proses ditentukan oleh status posting sesi',710)
box(d,(565,145,670,105),'Buka sesi → Edit Therapy Plan')
box(d,(80,330,740,125),'Sesi belum posted','Plan baru → infus & material draft diselaraskan')
box(d,(980,330,740,125),'Sesi sudah posted','Alasan + konfirmasi → reversal stok / finance')
arrow(d,(740,250),(450,330));arrow(d,(1060,250),(1350,330))
box(d,(980,540,740,125),'Pending → selaraskan → selesaikan ulang','Plan, infus aktual, material & audit terhubung')
arrow(d,(1350,455),(1350,540));d.text((80,560),'Kontrol transaksi',font=font(32,True),fill=NAVY)
d.text((80,610),'Stok kurang → seluruh edit dibatalkan',font=font(28),fill=GRAY)
save(im,'flow_koreksi')

im,d=canvas('Pembayaran dan kelayakan insentif','Status pembayaran dan verifikasi harus dibaca bersama',570)
for x,t,s in [(35,'Invoice','DRAFT → pending'),(625,'Pembayaran','Nominal + bukti'),(1215,'Verifikasi','VERIFIED / REJECTED')]:box(d,(x,155,545,150),t,s)
arrow(d,(580,230),(625,230));arrow(d,(1170,230),(1215,230))
box(d,(625,385,1135,125),'Invoice lunas & bukti terverifikasi','Data transaksi valid → perhitungan / review insentif')
arrow(d,(1487,305),(1487,385));d.text((35,395),'Bukti ditolak?',font=font(31,True),fill=NAVY)
d.text((35,445),'Perbaiki lalu unggah ulang.',font=font(28),fill=GRAY);save(im,'flow_payment')

im,d=canvas('Persediaan: pergerakan fisik dan pencatatan','Saldo ledger digunakan untuk pengendalian material dan valuasi',560)
for x,t,s in [(35,'Master & biaya','SKU • satuan • unit cost'),(625,'Stok masuk','GR • shipment • saldo awal'),(1215,'Pemakaian / jual','Sesi • add-on')]: box(d,(x,145,545,145),t,s)
arrow(d,(580,217),(625,217));arrow(d,(1170,217),(1215,217))
box(d,(625,375,1135,130),'Ledger → valuasi / HPP → jurnal & laporan','Pembatalan atau koreksi memakai pengembalian / reversal terkontrol')
arrow(d,(1487,290),(1487,375));save(im,'flow_inventory')

im,d=canvas('Snapshot kelengkapan pengisian','Kumulatif sampai ekspor 30 September 2026, 21.34 WIB • semua tanggal',550)
d.rounded_rectangle((40,160,1760,260),15,fill='#E7EDF2');d.rounded_rectangle((40,160,1417,260),15,fill=TEAL)
d.text((50,300),'3.143 lengkap',font=font(50,True),fill=TEAL);d.text((1190,300),'784 belum lengkap',font=font(45,True),fill=NAVY)
d.text((50,365),'80,04% dari 3.927 keterlibatan staff–sesi',font=font(31),fill=GRAY)
d.text((50,440),'29 / 39 staff memiliki sesi • 23 staff memiliki backlog',font=font(33,True),fill=NAVY);save(im,'chart_snapshot')

im,d=canvas('Tingkat kelengkapan berdasarkan role','Sumber: laporan bulanan September 2026 • snapshot kumulatif',650)
for i,(label,pct,backlog) in enumerate([('Dokter',70.53,402),('Nakes',85.32,174),('MSO',83.72,148),('Admin Cabang',87.21,60)]):
    y=160+i*110;d.text((35,y+15),label,font=font(35,True),fill=NAVY)
    d.rounded_rectangle((370,y,1450,y+60),10,fill=LIGHT);d.rounded_rectangle((370,y,370+1080*pct/100,y+60),10,fill=TEAL)
    d.text((1480,y),f'{pct:.2f}%'.replace('.',','),font=font(35,True),fill=NAVY)
    d.text((1480,y+43),f'Backlog: {backlog}',font=font(24),fill=GRAY)
save(im,'chart_roles')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492);sec.different_first_page_header_footer=True
styles=doc.styles
def style(name,size,color,before=0,after=6,line=1.25,bold=False):
    s=styles[name] if name in styles else styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(color.strip('#'));s.font.bold=bold
    f=s.paragraph_format;f.space_before=Pt(before);f.space_after=Pt(after);f.line_spacing=line
    return s
style('Normal',11,NAVY)
style('Title',32,NAVY,0,12,1.1,True)
style('Subtitle',15,GRAY,0,10,1.15)
for n,size,c,b,a in [(1,16,BLUE,18,10),(2,13,BLUE,14,7),(3,12,'#1F4D78',10,5)]:
    s=style(f'Heading {n}',size,c,b,a,1.15,True);s.paragraph_format.keep_with_next=True
style('Caption',9,GRAY,0,7,1.1)
style('Table Text',10,NAVY,0,4,1.1)
style('Source',8.5,GRAY,4,4,1.1)
style('Callout',10.5,NAVY,6,9,1.2)
style('Code',9,NAVY,0,5,1.1).font.name='Consolas'
style('Kicker',10,GOLD,0,8,1.1,True)
for name in ['List Bullet','List Number']:
    s=style(name,11,NAVY,0,4,1.25);s.paragraph_format.left_indent=Inches(.375);s.paragraph_format.first_line_indent=Inches(-.188)
numbering=doc.part.numbering_part.element
for n,fmt,text in [(80,'bullet','•'),(81,'decimal','%1.')]:
    ab=OxmlElement('w:abstractNum');ab.set(qn('w:abstractNumId'),str(n))
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,attrs in [('w:start',{'val':'1'}),('w:numFmt',{'val':fmt}),('w:lvlText',{'val':text}),('w:lvlJc',{'val':'left'})]:
        el=OxmlElement(tag)
        for k,v in attrs.items():el.set(qn('w:'+k),v)
        lvl.append(el)
    pp=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),'540');ind.set(qn('w:hanging'),'270');pp.append(ind)
    tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'num');tab.set(qn('w:pos'),'540');tabs.append(tab);pp.append(tabs);lvl.append(pp);ab.append(lvl);numbering.append(ab)
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(n));ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),str(n));num.append(ref);numbering.append(num)

hp=sec.header.paragraphs[0];hp.text='RAHO ERP  /  DOKUMENTASI PROJECT';hp.style='Source'
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('INTERNAL  •  Versi 1.1    |    Halaman ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
doc.core_properties.title='Dokumentasi Project RAHO ERP'
doc.core_properties.subject='Panduan sistem, operasional, arsitektur dan pemeliharaan'
doc.core_properties.author='RAHO';doc.core_properties.keywords='RAHO, ERP, klinik, dokumentasi'

def p(t,sty=None):return doc.add_paragraph(t,sty)
def h(t):return doc.add_heading(t,2)
def bullet(t):
    pp=p(t,'List Bullet');pr=pp._p.get_or_add_pPr();nu=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'80');nu.append(ni);pr.append(nu)
def steps(items):
    # Each procedure gets its own numbering instance, ensuring restart at 1.
    numid=100+len(numbering.findall(qn('w:num')));num=OxmlElement('w:num');num.set(qn('w:numId'),str(numid));ref=OxmlElement('w:abstractNumId');ref.set(qn('w:val'),'81');num.append(ref);ov=OxmlElement('w:lvlOverride');ov.set(qn('w:ilvl'),'0');st=OxmlElement('w:startOverride');st.set(qn('w:val'),'1');ov.append(st);num.append(ov);numbering.append(num)
    for t in items:
        pp=p(t,'List Number');pr=pp._p.get_or_add_pPr();nu=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),str(numid));nu.append(ni);pr.append(nu)
def note(t):
    pp=p(t,'Callout');shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F4F6F9');pp._p.get_or_add_pPr().append(shade)
def fig(name,cap,width=6.5):
    pp=doc.add_paragraph();pp.paragraph_format.keep_with_next=True;pp.paragraph_format.space_after=Pt(3)
    shape=pp.add_run().add_picture(str(BUILD/(name+'.png')),width=Inches(width));shape._inline.docPr.set('descr',cap)
    p(cap,'Caption')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    apply_table_geometry(t,widths)
    for i,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for pp in c.paragraphs:
                pp.style='Table Text'
                if i==0:
                    for r in pp.runs:r.bold=True
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF5');c._tc.get_or_add_tcPr().append(sh)
        if i==0:
            rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
    p('','Source');return t

chapters=[]
def chapter(n,title):
    doc.add_page_break();pp=doc.add_heading(f'{n:02d}  {title}',1);pp.paragraph_format.space_before=Pt(0)
    bm=OxmlElement('w:bookmarkStart');bm.set(qn('w:id'),str(n));bm.set(qn('w:name'),f'bab{n}');pp._p.insert(0,bm);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(n));pp._p.append(end);chapters.append(title)

# Cover: editorial_cover, named brand overrides (gold kicker, large title, logo).
pp=p('');pp.paragraph_format.space_after=Pt(32)
pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
logo=pp.add_run().add_picture(str(ROOT/'apps/web/public/asset/LOGORAHO.png'),width=Inches(2.45));logo._inline.docPr.set('descr','Logo resmi RAHO Club Premier dari repository')
pp=p('PROJECT DOCUMENTATION','Kicker');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_before=Pt(24)
pp=p('RAHO ERP','Title');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp=p('Clinic Management System','Subtitle');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp=p('Panduan sistem, proses operasional,\narsitektur dan pemeliharaan','Subtitle');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp=p('Member • Terapi • Inventory • Finance • Insentif','Caption');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(24)
fig('flow_utama','Peta proses terintegrasi RAHO. Diagram dokumentasi berdasarkan sumber project.')
pp=p('Versi 1.1  |  Disusun 12 Oktober 2026  |  Bahasa Indonesia','Source');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
pp=p('Baseline repository: 1dd9abd • 30 September 2026\nDokumen internal untuk manajemen, operasional dan tim IT','Source');pp.alignment=WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break();doc.add_heading('Panduan membaca dokumen',1)
p('Dokumentasi ini menghubungkan tujuan bisnis, pembagian kerja dan implementasi teknis RAHO. Pengguna operasional dapat membaca bab proses; tim IT dapat melanjutkan ke arsitektur, pengembangan dan deployment.')
nav=[('01','Gambaran project dan cakupan'),('02','Peta modul dan navigasi'),('03','Role, permission dan scope cabang'),('04','Alur operasional end-to-end'),('05','Member, paket dan dokumen'),('06','Pembayaran dan verifikasi'),('07','Sesi terapi dan finalisasi'),('08','Koreksi klinis dan reversal'),('09','Inventory, logistik dan add-on'),('10','Finance dan rekonsiliasi'),('11','Insentif bulanan staff'),('12','Voucher, WhatsApp, kolaborasi dan RAIN'),('13','Arsitektur dan struktur repository'),('14','Setup development dan konfigurasi'),('15','Deployment, backup dan recovery'),('16','QA, UAT dan penanganan masalah'),('17','Snapshot perkembangan dan penggunaan'),('18','Glosarium, sumber dan pemeliharaan')]
for n,t in nav:
    pp=p('', 'Normal');pp.paragraph_format.space_after=Pt(3)
    link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'bab'+str(int(n)));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'2E74B5');rp.append(co);r.append(rp);tx=OxmlElement('w:t');tx.text=f'{n}   {t}';r.append(tx);link.append(r);pp._p.append(link)
note('Cakupan bukti: kondisi source pada commit 1dd9abd. Informasi deployment dan hasil pengujian September berasal dari laporan project; bukan verifikasi ulang server pada 12 Oktober 2026. Diagram bukan screenshot aplikasi.')

chapter(1,'Gambaran project dan cakupan')
p('RAHO ERP adalah sistem manajemen klinik berbasis web untuk menghubungkan administrasi member, paket terapi, pelaksanaan tindakan, penggunaan material, stok, pembayaran dan pelaporan. Dukungan multi-cabang memungkinkan pengguna bekerja pada lingkup yang ditetapkan administrator.')
h('Tujuan sistem')
bullet('Menyatukan pencatatan klinis dan operasional agar satu transaksi dapat ditelusuri antar modul.')
bullet('Mengendalikan akses berdasarkan role, permission, cabang dan assignment sesi.')
bullet('Mengurangi selisih plan, infus aktual, material, persediaan serta data keuangan.')
bullet('Menyediakan laporan operasional, rekonsiliasi, audit dan perhitungan insentif.')
h('Audiens dan manfaat')
table(['Audiens','Kebutuhan utama'],[('Manajemen / Product Owner','Monitoring cabang, adopsi dan prioritas pengembangan'),('Operasional / tenaga medis','Panduan member, pembayaran dan pengisian sesi'),('Finance / Logistik','Kontrol tagihan, stok, valuasi dan rekonsiliasi'),('Tim IT','Arsitektur, konfigurasi, deployment dan dukungan')],[2800,6560])
h('Batas cakupan dokumentasi')
p('Dokumen menjelaskan kemampuan yang teridentifikasi dari repository dan panduan internal. Ketersediaan fitur pada suatu akun tetap bergantung pada permission, scope, konfigurasi dan versi deployment. Formula klinis serta kebijakan nominal insentif mengikuti keputusan organisasi, bukan ditetapkan oleh dokumen ini.')
note('Baseline: RAHO_APPS, commit 1dd9abd. Dokumen yang lebih lama digunakan sebagai konteks; package manifest, schema dan aturan workflow source menjadi rujukan ketika terdapat perbedaan.')

chapter(2,'Peta modul dan navigasi')
p('Mulai dari dashboard sesuai role, lalu pilih modul berdasarkan pekerjaan. Modul operasional saling terhubung; perubahan data final perlu memperhatikan dampaknya terhadap stok dan finance.')
table(['Area','Fungsi','Contoh halaman'],[
('Member & paket','Identitas, dokumen, kuota dan layanan','Member / detail member'),('Sesi terapi','Plan, tindakan, bahan dan evaluasi','/sessions'),('Pembayaran','Invoice, bukti dan verifikasi','/payments'),('Inventory','Ledger, biaya dan mutasi','Inventory / material usage'),('Finance','Jurnal, expense dan laporan','/accounting; /finance-reports'),('Kas / bank & purchasing','Rekening, supplier, PR, PO dan AP','/cash-bank; /purchasing'),('Staff incentives','Insentif bulanan per kelompok','/staff-incentives'),('Voucher','Campaign, registry dan klaim','/extra/vouchers'),('Kolaborasi & RAIN','Tim, task, subtask dan integrasi','/extra/collaboration; /extra/rain')],[2100,4100,3160])
h('Cara membaca navigasi')
p('Nama menu dapat berbeda sesuai role. Halaman yang dapat dibuka tidak selalu memberikan hak untuk mengubah data. Server memeriksa izin tindakan dan cabang, selain pembatasan yang terlihat pada UI.')
h('Navigasi tambahan September')
bullet('Insentif: /staff-incentives/nakes, /mso, /coordinator dan /doctor-head di bawah kelompok Staff Incentives.')
bullet('Voucher: campaigns, history, locations, operators dan registry di bawah /extra/vouchers.')
bullet('Kolaborasi: tasks dan teams di bawah /extra/collaboration; RAIN pada /extra/rain.')
p('Sumber: daftar fitur per role, panduan Finance, laporan bulanan September dan struktur halaman web.','Source')

chapter(3,'Role, permission dan scope cabang')
p('Role menentukan jenis tanggung jawab; permission menentukan tindakan yang diizinkan; scope cabang dan assignment membatasi objek yang dapat diakses. Ketiganya perlu diperiksa bersama ketika menu atau tombol tidak tersedia.')
table(['Peran','Tanggung jawab utama'],[('Super Admin','Konfigurasi, user, permission, audit dan pengawasan sistem'),('Admin Manager','Monitoring dan operasi pada beberapa cabang sesuai scope'),('Admin Cabang','Member, paket, sesi dan operasional satu cabang'),('Admin Layanan / MSO','Administrasi layanan, member, transaksi dan sesi'),('Admin Logistik','Persediaan, pengiriman, penerimaan dan purchasing'),('Finance / Controller','Keuangan, rekonsiliasi, approval dan kontrol logistik'),('Dokter','Diagnosis, therapy plan, infus sesuai izin dan evaluasi'),('Nakes / Perawat','Tanda vital, tindakan, material dan dokumentasi'),('Voucher Operator','Operasi voucher sesuai izin dan penugasan'),('Member','Akses data dan layanan pribadi')],[2850,6510])
h('Kontrol akses yang harus dipahami')
bullet('Pembatasan cabang diverifikasi backend pada operasi member dan sesi.')
bullet('Dokter yang ditugaskan dapat menyimpan infus aktual dan evaluasi secara independen; finalisasi sesi tetap memerlukan kelengkapan wajib.')
bullet('Laporan insentif seluruh staff tersedia bagi Super Admin dan Finance; role lain menerima data insentif dirinya sesuai aturan backend.')
note('Istilah Finance pada UI dapat terkait template permission. Schema baseline juga mengenal FINANCE_LOGISTICS_CONTROLLER. Jangan menyimpulkan kewenangan hanya dari label jabatan; periksa role dan permission aktif.')

chapter(4,'Alur operasional end-to-end')
fig('flow_utama','Gambar 1. Hubungan member, transaksi, terapi, material, finance dan laporan.')
steps(['Periksa akun dan cabang aktif sebelum memulai pekerjaan.','Cari member; daftarkan hanya jika belum ada. Lengkapi identitas dan dokumen yang diperlukan.','Tetapkan paket menggunakan katalog harga dan scope yang sesuai. Tinjau kuota, booster, diskon dan invoice.','Catat pembayaran dan verifikasi bukti sesuai kewenangan. Pastikan paket memenuhi syarat digunakan.','Buat sesi dan tetapkan staff yang bertanggung jawab. Lengkapi rencana serta catatan tindakan.','Periksa infus aktual, material, tanda vital dan evaluasi; finalisasi setelah validasi kelengkapan terpenuhi.','Periksa mutasi stok, penggunaan benefit, jurnal dan laporan. Tindak lanjuti selisih melalui koreksi resmi.'])
note('Output yang harus konsisten: data sesi, benefit paket, ledger persediaan, tagihan / pembayaran, jurnal dan audit. Integrasi eksternal mengikuti transaksi ERP yang valid.')

chapter(5,'Member, paket dan dokumen')
h('Registrasi dan pemeliharaan member')
steps(['Cari member menggunakan identitas yang tersedia untuk menghindari registrasi ganda.','Buka data member; periksa cabang, identitas dan riwayat yang relevan.','Lengkapi profil, foto serta dokumen pendukung atau informed consent sesuai proses yang berlaku.','Simpan satu kali dan tunggu hasil proses. Jika unggahan gagal, periksa status sebelum mencoba kembali.'])
h('Assignment dan perubahan paket')
p('Paket menghubungkan layanan, harga, kuota dan invoice. Katalog harus sesuai scope cabang. Super Admin dapat melakukan assignment lintas cabang sesuai kewenangan. Perubahan kuota perlu mempertahankan sesi yang telah digunakan.')
bullet('Tinjau total paket, booster, diskon dan nilai pembayaran sebelum konfirmasi.')
bullet('Status pembayaran dan status paket perlu dibaca bersama; invoice cicilan memerlukan penelusuran pembayaran yang tepat.')
bullet('Refund, pembatalan dan perubahan paket mengikuti status serta kontrol yang tersedia. Catat alasan tindakan penting.')
h('Dokumen dan penghapusan data')
p('Unggahan dokumen dilindungi terhadap submit berulang dan gangguan jaringan. Deaktivasi dan hard destruction member mempunyai guard terkait finance serta inventory; tindakan ini memerlukan penanganan oleh pengguna yang berwenang.')
note('Hasil yang diperiksa: member tersimpan pada scope yang benar, dokumen dapat dibuka oleh akun berizin, paket dan invoice sesuai transaksi, kuota terpakai tetap terlacak.')

chapter(6,'Pembayaran dan verifikasi')
fig('flow_payment','Gambar 2. Hubungan invoice, pembayaran, verifikasi dan kelayakan data insentif.')
h('Prosedur pembayaran')
steps(['Pilih invoice yang benar; periksa member, cabang, item, total dan sisa tagihan.','Catat nominal, metode dan akun kas / bank yang sesuai. Unggah bukti jika diwajibkan.','Pengguna berwenang memeriksa bukti lalu memverifikasi atau menolak dengan alasan.','Jika ditolak, perbaiki bukti atau data dan ajukan kembali. Hindari membuat pembayaran duplikat.','Pastikan nilai lunas, verification status dan status paket / invoice saling konsisten.'])
table(['Objek','Status penting','Interpretasi'],[('Invoice','DRAFT / PENDING_PAYMENT / PAID','Draft, tagihan aktif, atau lunas'),('Payment','PENDING / VERIFIED / REJECTED','Menunggu pemeriksaan, valid, atau ditolak')],[1900,3650,3810])
note('PAID tidak boleh dipakai sebagai satu-satunya indikator bukti sudah verified. Perbaikan September merekonsiliasi invoice lunas yang mempunyai bukti verifikasi agar tidak tertinggal PENDING. Nakes mendapat dukungan unggah dan konfirmasi pembayaran sesuai flow serta izin.')

chapter(7,'Sesi terapi dan finalisasi')
fig('flow_sesi','Gambar 3. Komponen pengisian sesi dan kontrol sebelum finalisasi.')
h('Pembagian kerja dan kelengkapan')
p('Dokter bertanggung jawab atas diagnosis, therapy plan dan evaluasi; Nakes mencatat kondisi, tindakan serta material aktual. Source UI memiliki penomoran langkah sampai 9, sedangkan kontrak StepCompletion menyimpan delapan flag. Diagram merangkum proses tanpa menyamakan nomor UI dan flag backend.')
bullet('Lengkapi diagnosis, plan, tanda vital sebelum, infus aktual, material, tanda vital sesudah dan evaluasi dokter sebagai komponen wajib yang terdaftar pada REQUIRED_SESSION_STEPS.')
bullet('Foto / dokumentasi disimpan sesuai kebutuhan proses; daftar required pada source baseline tidak memasukkan flag foto.')
bullet('Dokter dapat menyimpan infus dan evaluasi independen dari langkah awal; hal ini tidak melonggarkan validasi finalisasi.')
bullet('Reminder menampilkan langkah yang masih kurang sesuai tanggung jawab dan akses pengguna.')
note('Sebelum Selesaikan Sesi: periksa member, tanggal, assignment, sumber stok, infus aktual, material dan ringkasan kelengkapan. Sesi historis dengan skipInventoryConsumption tidak mengurangi persediaan saat ini.')
p('Sumber teknis: apps/web/src/components/sessions/sessionWorkflow.ts dan apps/web/src/types/session.ts.','Source')

chapter(8,'Koreksi klinis dan reversal')
fig('flow_koreksi','Gambar 4. Koreksi therapy plan terpakai berdasarkan status posting sesi.')
h('Prosedur koreksi')
steps(['Temukan sesi dan therapy plan terkait; periksa status posting dan dampak perubahan dosis.','Untuk sesi belum posted, gunakan edit yang tersedia; plan versi baru menyelaraskan infus serta material draft.','Untuk sesi posted, isi alasan dan konfirmasi reversal sesuai flow. Sistem mengembalikan sesi ke pending setelah koreksi stok / finance.','Periksa kembali plan, infus aktual dan material yang diselaraskan; lengkapi sesi dan selesaikan ulang.'])
h('Pengaman integritas')
bullet('Perubahan dilakukan dalam transaksi database; kekurangan stok membatalkan seluruh edit.')
bullet('Master inventory yang belum tersedia menghasilkan pesan yang perlu ditindaklanjuti.')
bullet('Material posted memerlukan reversal sebelum perubahan; edit keterangan tidak menimpa dosis aktual.')
bullet('Sesi historis yang bypass inventory tidak mengubah stok berjalan.')
note('Catat alasan koreksi dan review hasil akhir: jumlah material, mutasi ledger, dampak finance, status sesi serta audit. Reversal adalah koreksi terkontrol yang tetap menyimpan jejak transaksi.')

chapter(9,'Inventory, logistik dan add-on')
fig('flow_inventory','Gambar 5. Master produk, stok masuk, pemakaian dan pencatatan ledger.')
h('Prasyarat inventory')
bullet('Master produk, SKU, satuan dan lokasi / cabang harus tersedia untuk material yang dipakai.')
bullet('Default unit cost dan valuasi saldo lama perlu dilengkapi; dashboard membantu menemukan SKU tanpa biaya.')
bullet('Saldo ledger menjadi rujukan material infus untuk sesi baru; mirror inventory_items.stock dapat tertinggal.')
h('Alur logistik')
p('Permintaan stok diikuti proses persetujuan, pengiriman dan penerimaan sesuai scope. Purchasing menghubungkan supplier, PR, PO, Goods Receipt dan tagihan supplier. Inventori tim / tas serta pinjaman antartim mengikuti pengaturan cabang dan akses tim.')
h('Penjualan add-on dan Air Nano')
p('Penjualan harus memiliki nilai inventory yang valid. Pembatalan menggunakan pengembalian stok terkontrol. Pencarian SKU Air Nano mendukung variasi legacy; atribusi penjualan, jumlah dus dan status verifikasi pembayaran memengaruhi perhitungan insentif MSO.')
note('Jika stok kurang atau MATERIAL_INVENTORY_NOT_FOUND: periksa master, saldo ledger, satuan dan cabang. Jangan mengganti jumlah material hanya untuk melewati validasi. Koreksi data melalui proses yang berwenang.')

chapter(10,'Finance dan rekonsiliasi')
p('Finance menghubungkan invoice, pembayaran, kas / bank, persediaan, HPP, pengakuan pendapatan dan laporan. Data operasional ERP menjadi sumber pencatatan; jurnal serta integrasi mengikuti dokumen yang sudah valid.')
table(['Modul','Fungsi utama'],[('Accounting','Chart of accounts, periode dan jurnal'),('Kas & Bank','Rekening serta ledger penerimaan / pengeluaran'),('Opening Balance','Saldo awal pada cutover'),('Expense / Reimburse','Beban, bukti, approval dan pembayaran'),('Purchasing & AP','Supplier, PO, GR, bill dan pembayaran utang'),('Deferred Revenue','Pendapatan diterima di muka dan pengakuan per sesi'),('Finance Reports','P&L, Trial Balance, General Ledger dan rekonsiliasi')],[2900,6460])
h('Kontrol transaksi')
bullet('Periode harus OPEN; akun aktif dan diizinkan; debit / kredit seimbang.')
bullet('Permission dan scope tetap berlaku pada pembuatan, approval, pembayaran dan posting.')
bullet('Retry tidak boleh membuat jurnal atau kas / bank duplikat. Dokumen POSTED dikoreksi melalui reversal.')
h('Autonomous Finance dan approval')
p('Panduan Finance mendeskripsikan pengguna dengan template Finance aktif dapat membuat, memeriksa, memfinalisasi dan mem-posting dokumen miliknya jika berizin. Pengguna non-Finance tetap mengikuti maker-checker ketika diwajibkan modul. Karena itu, aturan self-approval harus dilihat per template dan proses.')
h('Rekonsiliasi yang disarankan')
p('Bandingkan invoice dengan pembayaran dan verifikasi, sesi dengan pengakuan pendapatan, ledger stok dengan material serta HPP, dan antrean Zoho dengan dokumen final ERP. Catat selisih, pemilik tindak lanjut dan hasil koreksi.')

chapter(11,'Insentif bulanan staff')
p('Staff Incentives menyediakan perhitungan bulanan yang menghubungkan data klinis, visit dan penjualan. Kebijakan tarif, target dan bonus harus mengikuti persetujuan Finance / HR.')
table(['Kelompok','Komponen yang didokumentasikan'],[('Nakes','Jumlah infus, tarif per infus, target dan bonus'),('MSO','Visit, Air Nano lunas, jumlah dus dan bonus target'),('Koordinator CHS','Assignment multi-cabang, infus personal dan target tim'),('Doctor Head','Assignment cabang dan target homecare / partnership')],[2750,6610])
h('Prosedur kerja bulanan')
steps(['Pilih bulan, cabang dan kelompok penerima sesuai akses.','Periksa assignment, jumlah infus / visit serta data penjualan yang memenuhi syarat.','Hitung atau perbarui draft menggunakan data terbaru; telusuri selisih sebelum review.','Lakukan review dan approval sesuai workflow. Snapshot reviewed / approved dipertahankan agar data yang disetujui terlindungi.','Finance / Super Admin melakukan koreksi terkontrol sebelum paid apabila diperlukan.','Ekspor laporan Excel sesuai scope; dokumentasikan penyelesaian pembayaran.'])
h('Akurasi dan privasi')
bullet('Air Nano dihitung berdasarkan kuantitas dus, bukan jumlah baris transaksi.')
bullet('Invoice PAID dengan bukti verifikasi direkonsiliasi; periksa kedua indikator sebelum memasukkan transaksi.')
bullet('Pembatasan laporan dan ekspor diterapkan backend. Role selain Finance / Super Admin hanya menerima insentif miliknya sesuai aturan.')
note('Dokumen ini tidak menetapkan nominal tarif atau target. Simpan kebijakan bisnis yang disetujui beserta masa berlakunya sebagai rujukan perhitungan.')

chapter(12,'Voucher, WhatsApp, kolaborasi dan RAIN')
h('Voucher dan partnership')
p('Kelompok Voucher mencakup dashboard, campaign, registry, operator, lokasi dan histori klaim. Campaign mengatur penerbitan serta ketersediaan voucher; registry menghubungkan kode dan status dengan proses klaim. Lokasi klaim mengikuti kebijakan campaign, termasuk lokasi reguler yang didukung.')
steps(['Siapkan campaign dan kebijakan lokasi melalui akun berizin.','Terbitkan atau kelola voucher pada registry; periksa masa berlaku dan status ketersediaan.','Operator melakukan klaim sesuai lokasi dan ketentuan; periksa histori untuk penelusuran.'])
note('Voucher dummy adalah sarana pengujian; pisahkan penggunaannya dari laporan operasional produksi.')
h('WhatsApp dan laporan sesi')
p('Pairing menggunakan QR code, dengan timeout yang diperpanjang. Template laporan mendukung background dan branding khusus. Foto sesi dapat disertakan sesuai flow laporan serta akses data pengguna.')
bullet('Lakukan pairing dari menu yang berizin; periksa koneksi sebelum pengiriman.')
bullet('Tinjau isi, tujuan dan gambar laporan untuk memastikan data sesi yang dikirim benar.')
h('Kolaborasi tim')
p('Area /extra/collaboration menyediakan team, membership, task, subtask, owner utama dan monitoring progres. Histori assignment membantu menelusuri perubahan tanggung jawab. Pemilik dan leader perlu disesuaikan dengan proses organisasi.')
h('RAIN / OpenClaw')
p('RAIN mempunyai halaman UI, personal task API dan integrasi awal OpenClaw yang terautentikasi. Lakukan UAT role, pembatasan akses dan monitoring error sebelum penggunaan organisasi diperluas. Status integrasi tidak boleh diasumsikan aktif hanya karena halaman tersedia.')

chapter(13,'Arsitektur dan struktur repository')
fig('arsitektur','Gambar 6. Arsitektur komponen berdasarkan manifest, schema dan konfigurasi project.')
table(['Lokasi','Isi / tanggung jawab'],[('apps/web','Next.js 14.2.35, React 18, TypeScript, halaman dan komponen'),('apps/api/src','Express, controller, service, middleware dan modul'),('apps/api/prisma','Schema, migration dan seed PostgreSQL'),('docs','Panduan, requirement, UAT dan laporan perkembangan'),('integrations/openclaw','Workspace dan plugin integrasi RAIN / OpenClaw'),('docker-compose.*.yml','Konfigurasi infrastruktur development / production')],[3100,6260])
p('Data bisnis disimpan melalui Prisma pada PostgreSQL. MinIO menyediakan object storage. JWT digunakan untuk autentikasi; validasi, permission dan branch scope dijalankan di API. Web menggunakan Axios, Zustand, React Hook Form dan dependensi UI sesuai manifest.','Source')
doc.add_page_break()
h('13.1 Tech stack aplikasi')
p('RAHO ERP menggunakan arsitektur full-stack berbasis TypeScript dalam satu monorepo. Frontend mengelola interaksi pengguna, backend menjalankan aturan bisnis dan kontrol akses, sedangkan database serta object storage menyimpan data dan file.')
table(['Bagian','Teknologi','Fungsi'],[
('Frontend','Next.js 14.2.35; React 18; TypeScript','Halaman staff / member dan komponen aplikasi'),
('Styling','Tailwind CSS; CSS Modules; UI khusus','Layout, tema dan tampilan komponen'),
('State','Zustand; React Hooks','State antarmuka dan interaksi pengguna'),
('Form / validasi','React Hook Form; Zod','Pengelolaan input dan validasi data'),
('HTTP client','Axios','Komunikasi frontend dengan API'),
('Backend','Node.js; Express 4; TypeScript','REST API dan aturan bisnis klinik'),
('Database','PostgreSQL 15','Data member, sesi, stok dan transaksi'),
('ORM / migration','Prisma 5.22','Akses database, schema dan migration'),
('Autentikasi','JWT; bcryptjs','Token autentikasi dan hashing password'),
('Otorisasi','Role; permission; scope; assignment','Kontrol akses data dan tindakan'),
('File storage','MinIO; AWS SDK S3','Foto, dokumen dan bukti pembayaran')
],[1850,3650,3860])
note('Versi di atas merujuk manifest dan konfigurasi baseline 1dd9abd. Manifest root menetapkan Node.js >=18; versi runtime aktual mengikuti environment deployment dan kompatibilitas dependency lockfile.')
doc.add_page_break()
h('13.2 Pelaporan, integrasi dan tooling')
table(['Bagian','Teknologi','Fungsi'],[
('Grafik','Recharts','Visualisasi data pada halaman aplikasi'),
('Ekspor Excel','ExcelJS','Workbook dan laporan Excel'),
('Dokumen PDF','jsPDF; jsPDF AutoTable; PDFKit','Pembuatan PDF di frontend / backend'),
('WhatsApp','Baileys','Pairing dan pengiriman laporan sesi'),
('Finance eksternal','Zoho Books','Sinkronisasi dokumen bisnis sesuai konfigurasi'),
('Integrasi internal','OpenClaw / RAIN','Personal task API dan integrasi pendukung'),
('Unit / komponen','Jest; React Testing Library','Pengujian API dan komponen web'),
('End-to-end','Playwright','Pengujian alur pengguna melalui browser'),
('Kualitas kode','TypeScript; ESLint','Pemeriksaan tipe dan konsistensi kode'),
('Infrastruktur','Docker; Docker Compose','Container dan layanan development / production'),
('CI / deployment','GitHub Actions','Otomatisasi build, pengujian dan deployment')
],[1850,3650,3860])
h('Hubungan antar teknologi')
p('Browser pengguna membuka Next.js / React. Web mengirim permintaan ke Express API. API memeriksa autentikasi, validasi, permission dan scope, lalu mengakses PostgreSQL melalui Prisma. File diproses melalui MinIO, sedangkan WhatsApp, Zoho Books dan RAIN digunakan sesuai konfigurasi integrasi.')
note('Integrasi yang tersedia di source belum tentu aktif di production. Tailwind dan CSS Modules tercantum sebagai teknologi styling; keduanya dapat dipakai pada komponen berbeda. Versi dependency dengan rentang semver pada manifest bukan bukti versi runtime yang sedang berjalan.')
p('Rujukan: package.json; apps/web/package.json; apps/api/package.json; apps/api/prisma/schema.prisma; docker-compose.dev.yml; workflow repository.','Source')

chapter(14,'Setup development dan konfigurasi')
p('Gunakan environment development terpisah. Manifest root mensyaratkan Node.js >=18; versi Node yang dipilih perlu kompatibel dengan dependency lockfile. Docker development menyediakan PostgreSQL 15 dan MinIO.')
steps(['Pasang dependency dari root repository sesuai lockfile menggunakan npm ci.','Siapkan environment API dan Web berdasarkan contoh konfigurasi repository.','Jalankan infrastruktur development: docker compose -f docker-compose.dev.yml up -d.','Generate Prisma client: npm run db:generate --prefix apps/api.','Jalankan migrasi pada database development yang telah diverifikasi. Script predev API juga menjalankan deploy-with-recovery.cjs.','Jika membutuhkan data demo, pilih varian seed yang sesuai pada database development.','Jalankan npm run dev untuk API dan Web; periksa log dan koneksi database / storage.'])
h('Kelompok konfigurasi')
table(['Area','Yang perlu disiapkan'],[('Database','DATABASE_URL dengan host dan database yang tepat'),('Autentikasi','Konfigurasi JWT dan masa berlaku sesi sesuai environment'),('Storage','Endpoint, akses bucket dan konfigurasi MinIO'),('Web / API','Alamat API, origin dan konfigurasi lintas layanan'),('Integrasi','Zoho, WhatsApp serta RAIN hanya jika diperlukan')],[2650,6710])
h('Perintah validasi')
for t in ['npm run type-check','npm run build','npm run test:api','npm run test:web','npm run e2e:web']:p(t,'Code')
note('Nilai password, token dan connection string tidak dicantumkan pada dokumentasi distribusi. Seed / reset hanya untuk database development yang benar-benar ditujukan bagi pengujian.')

chapter(15,'Deployment, backup dan recovery')
h('Urutan release')
steps(['Tetapkan commit release dan periksa perubahan schema, konfigurasi serta integrasi.','Jalankan pemeriksaan tipe, build dan pengujian relevan; lakukan regression / UAT sesuai dampak.','Pastikan backup database dan object storage tersedia, lalu uji restore pada salinan environment.','Jalankan migrasi production secara berurutan melalui jalur deployment yang didokumentasikan.','Deploy API / Web dan periksa hasil workflow serta status layanan.','Lakukan smoke test login, member, sesi, inventory, pembayaran dan laporan insentif.','Catat commit, waktu deployment, hasil smoke test dan tindak lanjut.'])
h('Recovery migration legacy')
p('Helper deploy-with-recovery mengenali kegagalan tertentu pada migrasi legacy dan memverifikasi SQL sebelum melanjutkan migrate deploy. Recovery tidak mengizinkan bypass seluruh error. Kesalahan permission, integritas data atau SQL yang tidak dikenali memerlukan pemeriksaan manual oleh tim IT.')
bullet('Jangan menggunakan migrate reset untuk memperbaiki P3009 pada database berisi data.')
bullet('Recovery mandiri tidak otomatis membuat backup; pastikan backup tersedia lebih dulu.')
bullet('Jika recovery menolak, hentikan deployment, simpan log dan periksa database; jangan memperluas whitelist hanya agar deployment lewat.')
h('Evidence deployment September')
p('Laporan September mencatat commit 1dd9abd selesai deployed pada 30 September 2026 pukul 17.05 WIB melalui workflow Deploy RAHO_APPS, run 36699196578. Pernyataan ini merupakan evidence historis dari laporan, bukan pengecekan kondisi server saat dokumen dibuat.')
p('Rujukan: docs/DEPLOYMENT_MIGRATION_RECOVERY.md; workflow dan skrip backup / restore pada apps/api.','Source')

chapter(16,'QA, UAT dan penanganan masalah')
h('Skenario UAT lintas role')
bullet('MSO / Admin: cari member, assign paket, buat invoice dan unggah bukti pembayaran.')
bullet('Finance / role berizin: verifikasi bukti, periksa status paket dan dampak kas / bank.')
bullet('Dokter / Nakes: isi sesi, infus aktual, material, evaluasi dan finalisasi.')
bullet('Logistik / Finance: cocokkan material dengan ledger stok, HPP dan jurnal.')
bullet('Super Admin / Manager: uji koreksi sesi posted, reversal serta pembatasan cabang.')
bullet('Insentif: uji draft, snapshot, ekspor dan privasi data pribadi.')
h('Diagnosis masalah umum')
table(['Gejala','Pemeriksaan pertama'],[('Menu / tombol tidak tampil','Role, permission, template dan scope cabang'),('Sesi tidak bisa selesai','Reminder dan komponen wajib yang belum lengkap'),('Material tidak ditemukan','SKU, master inventory, lokasi dan cabang'),('Stok tidak cukup','Saldo ledger, sumber stok serta satuan material'),('Invoice lunas masih pending','Payment, bukti, verification status dan rekonsiliasi'),('Insentif berbeda','Periode, assignment, jumlah dus dan snapshot'),('Unggahan gagal','Koneksi, status unggahan dan izin storage'),('Integrasi gagal','Koneksi, autentikasi, scope dan log antrean')],[3200,6160])
h('Format eskalasi')
p('Sertakan modul, role, cabang, ID / kode dokumen, waktu WIB, langkah reproduksi, pesan error dan hasil yang diharapkan. Lampirkan screenshot seperlunya dengan data sensitif ditutup. Jangan mengulang transaksi sebelum memeriksa apakah proses sebelumnya sudah tersimpan.')
note('Bukti testing September: targeted API 28/28, Web 7/7, type-check dan build lulus menurut laporan. Ini bukan klaim seluruh test telah dijalankan ulang untuk penyusunan dokumen ini.')

chapter(17,'Snapshot perkembangan dan penggunaan')
fig('chart_snapshot','Gambar 7. Snapshot kumulatif keterlibatan staff–sesi, bukan aktivitas khusus September.')
fig('chart_roles','Gambar 8. Kelengkapan per role dan jumlah backlog pada saat ekspor.')
p('Sumber penggunaan: kinerja-staff-2026-09-30.xlsx, filter semua cabang, semua tanggal dan semua staff. Angka dikutip dari laporan September; file ekspor tidak dihitung ulang untuk dokumen ini. Satu sesi dapat melibatkan beberapa staff sehingga 3.927 bukan jumlah sesi unik global.','Source')
h('Tindak lanjut operasional')
bullet('Prioritaskan 784 backlog, terutama Dokter dan Raho Premier Jakarta; dokumentasikan penyebab sebelum menilai kinerja.')
bullet('Pisahkan cabang tes / dummy dari KPI dan gunakan filter tanggal bulanan untuk perbandingan periode.')
bullet('Tambahkan login terakhir, hari aktif serta modul yang dipakai untuk mengukur adopsi di luar sesi.')
p('Ringkasan engineering September menurut laporan: 113 commit, 436 file berubah, 27 migration baru, 122 file test berubah dan 15 halaman staff baru. Angka ini adalah perubahan terhadap baseline Agustus, bukan ukuran seluruh repository.','Source')

chapter(18,'Glosarium, sumber dan pemeliharaan')
table(['Istilah','Makna dalam dokumentasi'],[('Scope cabang','Lingkup cabang yang diizinkan untuk akun / tindakan'),('Assignment','Penugasan staff pada sesi, tim atau proses'),('Therapy plan','Rencana terapi yang menjadi rujukan tindakan'),('Infus aktual','Pencatatan pelaksanaan infus pada sesi'),('Ledger','Catatan mutasi yang membentuk saldo dan jejak transaksi'),('HPP','Harga pokok yang terkait penggunaan / penjualan persediaan'),('Posted','Dokumen telah diproses pada pencatatan final terkait'),('Reversal','Pembalikan transaksi terkontrol untuk koreksi'),('Snapshot insentif','Data perhitungan yang dipertahankan setelah review / approval'),('UAT','Uji penerimaan pengguna pada alur operasional')],[2700,6660])
h('Sumber utama di repository')
for t in ['README.md; package.json; apps/web/package.json; apps/api/package.json','apps/api/prisma/schema.prisma; docker-compose.dev.yml','docs/DAFTAR_FITUR_PER_ROLE_DAN_FLOW.md','docs/PANDUAN_LENGKAP_ERP_RAHO_UNTUK_PENGGUNA_AWAM.md','docs/PANDUAN_FITUR_FINANCE_DAN_FLOW.md','docs/DEPLOYMENT_MIGRATION_RECOVERY.md; docs/UAT/','docs/LAPORAN_BULANAN_SEPTEMBER_2026.md','apps/web/src/components/sessions/sessionWorkflow.ts']:p(t,'Source')
h('Pemeliharaan dokumen')
p('Perbarui dokumentasi ketika menu, permission, status, formula insentif atau jalur deployment berubah. Sertakan commit rujukan dan tanggal revisi; diagram mengikuti perubahan proses. Screenshot aplikasi dapat ditambahkan setelah tersedia environment demo yang representatif.')
p('Riwayat dokumen: v1.0 • 12 Oktober 2026 • dokumentasi terpadu baseline 1dd9abd. Revisi v1.1 • 12 Oktober 2026 • penambahan penjelasan tech stack.','Source')

path=OUT/'DOKUMENTASI_PROJECT_RAHO_ERP_v1.1.docx';doc.save(path)
issues=audit_docx_tables(path)
report={'preset':'compact_reference_guide','header':'editorial_cover','named_overrides':['brand cover gold kicker','32pt cover title','10pt table text','9pt captions','1.15 heading line spacing'],'chapters':len(chapters),'tables':len(doc.tables),'inline_images':len(doc.inline_shapes),'paragraphs':len(doc.paragraphs),'table_geometry_issues':issues,'baseline':'1dd9abd','output':str(path)}
(BUILD/'structural_qa.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps(report,indent=2))
